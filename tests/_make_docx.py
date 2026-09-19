from docx import Document
doc = Document()
doc.add_heading('Ali Murtaza', level=1)
doc.add_paragraph('Email: ali@example.com')
doc.add_paragraph('Phone: +1-555-1234')
doc.add_paragraph('Skills: Python, FastAPI, SQLAlchemy, Docker')
doc.add_paragraph('Experience: 5 years in backend development')
doc.save(r'E:\Ai Resume Analyzer\AI-Resume-Analyzer\tests\tmp_resume.docx')
print('created')
